"""Generate the final project report (docs/Scheme_Eligibility_Assistant_Report_Final.docx).

Dataset and extraction numbers are computed live from the project; test /
benchmark / evaluation numbers are read from docs/results.json, which
scripts/collect_results.py writes after a real run.
"""
import json
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).absolute().parent.parent
sys.path.insert(0, str(ROOT))
from api.engine.validate import leaves, validate_tree  # noqa: E402
from api.ingest.extract import propose  # noqa: E402

IMG = ROOT / "docs" / "img"
OUT = ROOT / "docs" / "Scheme_Eligibility_Assistant_Report_Final.docx"
ACCENT = RGBColor(0x1F, 0x3A, 0x8A)

# ---------------------------------------------------------------- live numbers
schemes = json.loads((ROOT / "data" / "schemes_parsed.json").read_text(encoding="utf-8"))
curated = json.loads((ROOT / "data" / "curated_rules.json").read_text(encoding="utf-8"))["schemes"]
n_pdfs = len(list((ROOT / "data" / "raw_pdfs").glob("*.pdf")))
levels = Counter(s["level"] for s in schemes)
n_states = len({s["state"] for s in schemes if s["state"]})
valid, attr_counts, leaf_total = 0, Counter(), 0
for s in schemes:
    p = propose(s)
    ls = leaves(p["eligibility"])
    leaf_total += len(ls)
    attr_counts.update(l["attribute"] for l in ls)
    if not validate_tree(p["eligibility"], " ".join(s["sections"].values())):
        valid += 1
results_path = ROOT / "docs" / "results.json"
R = json.loads(results_path.read_text(encoding="utf-8")) if results_path.exists() else {}


# ---------------------------------------------------------------- helpers
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.3))
sec.top_margin = sec.bottom_margin = Cm(2.2)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.15
for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
    st = styles[f"Heading {lvl}"]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = ACCENT
    st.font.bold = True
    st.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
    st.paragraph_format.space_after = Pt(6)


def para(text="", bold=False, italic=False, size=None, align=None, color=None):
    p = doc.add_paragraph()
    _runs(p, text, bold, italic, size, color)
    if align:
        p.alignment = align
    return p


def _runs(p, text, bold=False, italic=False, size=None, color=None):
    # **bold** spans inside text
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        _runs(p, it)
        p.paragraph_format.space_after = Pt(3)


def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(8)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F3F4F6")
    p._p.get_or_add_pPr().append(shd)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(8.5)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(header, rows, widths=None, size=9.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F3A8A")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _runs(cells[i].paragraphs[0], str(v), size=size)
            if ri % 2:
                shade(cells[i], "EEF2FF")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(path, caption, width=16):
    if Path(path).exists():
        doc.add_picture(str(path), width=Cm(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para(caption, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x55, 0x55, 0x55))


def figure_grid(paths, captions, width=7.9):
    """Screenshots two per row in a borderless table, each with its caption."""
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(0, len(paths), 2):
        cells = t.add_row().cells
        for j in range(2):
            if i + j >= len(paths):
                continue
            c = cells[j]
            c.width = Cm(8.2)
            c.paragraphs[0].add_run().add_picture(str(paths[i + j]), width=Cm(width))
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = c.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(cap, captions[i + j], italic=True, size=8.5, color=RGBColor(0x55, 0x55, 0x55))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def r(key, default="—"):
    cur = R
    for k in key.split("."):
        if not isinstance(cur, dict) or k not in cur:
            return default
        cur = cur[k]
    return cur


# ---------------------------------------------------------------- title page
GREY = RGBColor(0x6B, 0x72, 0x80)
CENTER = WD_ALIGN_PARAGRAPH.CENTER


def running_text(container, text):
    p = container.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if container is sec.header else CENTER
    run = p.add_run(text)
    run.font.size, run.font.color.rgb = Pt(8), GREY


running_text(sec.header, "Scheme Eligibility Assistant — Explainable Eligibility & Grounded Q&A | Technical Documentation")
running_text(sec.footer, "Scheme Eligibility Assistant Project Report — MongoDB Architecture & Rules Engine")


def tp(text, size=11, bold=False, color=None, space=6, underline=False):
    p = para(text, bold=bold, size=size, align=CENTER, color=color)
    p.paragraph_format.space_after = Pt(space)
    if underline:
        for run in p.runs:
            run.underline = True
    return p


tp("A NoSQL Report on", size=11, underline=True, space=10)
tp("Scheme Eligibility Assistant: Explainable Government-Scheme Eligibility & Grounded Q&A on MongoDB",
   size=20, bold=True, color=ACCENT, space=14)
tp("Submitted to Manipal University Jaipur", size=11, space=4)
tp("Towards the partial fulfillment for the Award of the Degree of", size=11, space=10)
tp("B. Tech Computer Science and Engineering (Artificial Intelligence and Machine Learning)", size=13, bold=True,
   color=ACCENT, space=8)
tp("In Computer Science and Engineering (AIML)", size=11, space=4)
tp("2024-2028", size=11, space=10)
tp("By", size=11, underline=True, space=6)
tp("Nalin Ahuja (2427010036)", size=11, space=4)
tp("Adamya Singh Thakur (2427010181)", size=11, space=16)
doc.add_picture(str(IMG / "muj_logo_transparent.png"), width=Cm(9.5))
doc.paragraphs[-1].alignment = CENTER
doc.paragraphs[-1].paragraph_format.space_after = Pt(22)
tp("Under the guidance of", size=11, space=4)
tp("Ms. Surabhi Agrawal", size=11, bold=True, space=16)
tp("Department of Artificial Intelligence and Machine Learning", size=15, bold=True, color=ACCENT, space=8)
tp("Manipal University Jaipur", size=15, bold=True, color=ACCENT, space=8)
tp("Jaipur, Rajasthan", size=15, bold=True, color=ACCENT, space=0)
page_break()

# ---------------------------------------------------------------- abstract + contents
doc.add_heading("Abstract", 1)
para(f"Every year the central and state governments announce a large number of welfare schemes, yet many people "
     f"who could benefit from them never apply. In most cases they simply do not know that they qualify. The rules "
     f"are written in long pages of official text, each scheme phrases them differently, and they keep getting "
     f"revised. We built the **Scheme Eligibility Assistant** to close this gap. It is a web application that "
     f"works in two ways. First, a user fills in a short profile and our rules engine goes through every published "
     f"scheme and tells them whether they are eligible, not eligible, or whether we still need some detail from "
     f"them — and for each scheme it shows which condition decided the answer. Second, the user can type a question "
     f"in plain English and get an answer that is put together only from the scheme text we have stored, with a "
     f"link back to the source. If our data does not cover the question, the assistant says so instead of making "
     f"something up. Signed-in citizens also get a ranked list of recommended schemes with the reasons behind each "
     f"rank, an application tracker with a document checklist and deadline reminders, and a Grievance Portal for "
     f"complaints. Citizens and officers sign in on separate tabs of a login page, and officers get their own "
     f"dashboard, which includes a coverage-gap analysis that maps where people search for help that the scheme data "
     f"does not cover.")
para(f"We used MongoDB for all of the storage. A scheme and its eligibility conditions live together in one "
     f"document. Whenever the conditions change we save a fresh version instead of editing the old one, and this "
     f"save happens inside a transaction. Every check a user runs is recorded with a copy of the exact rules and "
     f"answers that were used, plus a hash so that tampering can be detected later. For question answering we "
     f"combine MongoDB text search with similarity search over stored embeddings. Our data comes from "
     f"{n_pdfs} scheme pages published on myScheme (Hugging Face dataset shrijayan/gov_myscheme), which we cleaned "
     f"into {len(schemes)} schemes — {levels['central']} central and {levels['state']} from {n_states} states and union "
     f"territories. One idea guided the whole design: eligibility is always decided by plain, testable code, never "
     f"by a language model.")
para("Keywords: MongoDB, document database, schema validation, aggregation, $facet, $bucket, geospatial index, "
     "transactions, TTL index, full-text and vector search, rules engine, recommendation ranking.", italic=True, size=10)

doc.add_heading("Contents", 1)
for line in ["1. Problem identification and real-world relevance", "2. Objectives and scope",
             "3. NoSQL database selection and justification", "4. System architecture",
             "5. Database design and data modelling", "6. CRUD implementation",
             "7. Advanced NoSQL features and query optimisation", "8. Rules engine design",
             "9. Grounded retrieval-augmented Q&A", "10. Application integration and functionality",
             "11. Testing, security and data validation", "12. Results", "13. Limitations and future work",
             "14. Conclusion", "15. References", "Appendix A — Setup and run instructions",
             "Appendix B — Plagiarism report"]:
    _c = para(line, size=10.5)
    _c.paragraph_format.space_after = Pt(0)
    _c.paragraph_format.line_spacing = 1.0
page_break()

# ---------------------------------------------------------------- 1. Problem
doc.add_heading("1. Problem identification and real-world relevance", 1)
doc.add_heading("1.1 Problem statement", 2)
para("Suppose a farmer in Assam wants to know which government schemes she can apply for. Today she would have to "
     "open scheme pages one by one and work out conditions such as age limits, income ceilings, caste category, land "
     "owned, BPL status and state of residence — and then do the same for the next scheme, and the next. Most people "
     "do not have the time for this, so they either give up or pay an agent. Staff at Common Service Centres and "
     "block offices run into the same problem, only with many more applicants. A general chatbot is not a good "
     "answer either: if it confidently tells someone the wrong thing about a benefit, that person can lose money "
     "they were entitled to.")
doc.add_heading("1.2 Who it serves", 2)
bullets(["**Citizens**, who want to find out what they can claim without handing over identity documents.",
         "**CSC operators and block officers**, who screen many applicants a day and need answers they can explain.",
         "**Scheme administrators**, who need to change rules safely and look back at decisions made earlier."])
doc.add_heading("1.3 Why it matters", 2)
bullets(["A lot of welfare money goes unclaimed because the people it is meant for never hear that they are eligible.",
         "A useful answer has to be **explainable** — the person should see which condition they met or missed — and "
         "**reproducible**, meaning we can show what the rules were on the day they checked.",
         "Scheme rules are revised from time to time. Showing when our data was last verified, and keeping old rule "
         "versions instead of overwriting them, are practical requirements for a system like this."])

# ---------------------------------------------------------------- 2. Objectives
doc.add_heading("2. Objectives and scope", 1)
doc.add_heading("2.1 Objectives", 2)
bullets(["Store scheme data whose structure varies (central and state schemes are not laid out the same way) in a "
         "document database.",
         "Write eligibility conditions as data — a tree of rules — so that a single piece of code can check all the "
         "schemes, instead of writing separate code for each one.",
         "Give one of three answers for every scheme (Eligible, Not eligible, Need more information), explain each "
         "condition, and suggest the one question that would help the user most.",
         "Keep a record of every check that cannot be quietly edited and that can be re-run later from what it stored.",
         "Never edit published rules in place; publish each change as a new version in a single atomic step.",
         "Answer free-text questions using only stored scheme text, show the sources, and refuse when the sources do "
         "not contain the answer.",
         "Give users a way to raise a complaint about an application or a result, track it by ticket number, and see "
         "every update from the officer handling it.",
         "Rank the schemes that fit a person and explain the rank; let citizens save schemes and track their own "
         "applications; show officers where demand for schemes is not being met.",
         "Cover the database topics of the course: CRUD, indexing, aggregation, validation, TTL, text and vector "
         "search, login and roles, and automated tests."])
doc.add_heading("2.2 Scope", 2)
table(["Included in this project", "Left for future work"], [
    [f"{len(schemes)} schemes from the myScheme dataset; rules for {len(curated)} of them checked by hand",
     "Live updates from the official API Setu / myScheme services"],
    ["Eligibility check, ranked recommendations, Q&A, scheme browser, tracker, grievances; officer dashboard", "Applying to a scheme on the user's behalf"],
    ["English text (a multilingual embedding model can be swapped in later)", "A full Hindi or regional-language interface"],
    ["MongoDB running locally as a replica set; the code also supports Atlas", "Cloud hosting and production monitoring"],
], widths=[8.2, 8.2])
para("Note: this is a college project built on public data. It is not an official government service. Every scheme "
     "in the app links to its page on myScheme so that users can confirm the details there.", italic=True, size=10)

# ---------------------------------------------------------------- 3. NoSQL selection
doc.add_heading("3. NoSQL database selection and justification", 1)
doc.add_heading("3.1 Why a document database", 2)
para("When we looked at the data, a scheme turned out to be one nested object: a few text sections, a list of "
     "benefits, a list of documents to submit, and an eligibility condition that can be nested to any depth — for "
     "example A AND (B OR C) AND NOT D. Central and state schemes do not fit the same set of columns. If we had used "
     "a relational database, we would have needed separate tables for schemes, conditions, condition groups, benefits "
     "and documents, and joined them back together on every request. The nested condition tree would still have "
     "needed either a self-referencing table queried recursively or a JSON column. A document database lets us keep "
     "the whole thing together and read it in one go.")
table(["What we needed", "With a relational database", "With MongoDB (our choice)"], [
    ["Schemes with different fields", "Many empty columns, or key-value tables", "Flexible documents, checked by a JSON Schema validator"],
    ["Nested eligibility conditions", "Self-referencing table with recursive queries, or JSONB", "Stored as a nested sub-document and read in one fetch"],
    ["Load a scheme with its rules", "Several joins", "One document read"],
    ["Audit record with the exact inputs", "Joins back to versioned tables that must never change", "Copies embedded in the record itself"],
    ["Text plus similarity search", "Extra extensions", "Text index and vector search next to the data"],
    ["Publish rules atomically", "Transactions", "Multi-document transactions on a replica set"],
    ["Growing audit log", "Partitioning and read replicas", "Sharding on user or run keys; replica sets"],
], widths=[4.6, 5.6, 6.2], size=9)
doc.add_heading("3.2 Why MongoDB in particular", 2)
bullets(["**Validators on collections.** We keep the flexibility of documents, but the server still refuses bad data "
         "such as an unknown rule operator, an invalid verdict, a negative age or a missing checksum.",
         "**Aggregation pipelines.** Reports such as how many checks ended in each verdict, or which answers were "
         "missing most often, are computed inside the database.",
         "**Transactions.** On a replica set we can save a new rule version and point the scheme to it as one step.",
         "**TTL indexes.** Chat logs and rate-limit counters delete themselves after a set time, so we do not keep "
         "data longer than needed.",
         "**Search.** A text index handles keyword matching. On Atlas the code uses $vectorSearch for similarity; on "
         "our local server it computes the same similarity from embeddings stored in MongoDB.",
         "**A good Python driver.** PyMongo works well with our FastAPI backend."])
doc.add_heading("3.3 What we gave up", 2)
para("MongoDB does not have foreign keys, so we had to protect the data in other ways. We added unique indexes "
     "(scheme slug, slug plus version, email, one profile per user), validators on the important collections, "
     "Pydantic checks on every API request, and our own rule validator that runs before anything is published. We "
     "also follow a simple rule in the code: published versions and audit records are only ever inserted, never "
     "updated.")

# ---------------------------------------------------------------- 4. Architecture
doc.add_heading("4. System architecture", 1)
para("The system has three parts that all use the same MongoDB database. The ingestion part reads the scheme PDFs "
     "and prepares the data offline. The rules engine decides eligibility. The Q&A part finds relevant text and "
     "writes answers, and whenever a question is really about eligibility it passes the work to the rules engine. "
     "The recommendation ranking, the application tracker, the grievance portal and the coverage analytics use the "
     "same database. The browser only ever talks to our FastAPI server, through three pages: the public site, the "
     "login page and the officers' dashboard.")
figure(IMG / "architecture.png", "Figure 1 — System architecture")
table(["Layer", "What we used", "Why"], [
    ["Frontend", "Three pages (public site, /login, /admin) in HTML and plain JavaScript; Leaflet for the map", "No build step, supports streamed answers; the map needs internet for tiles"],
    ["Backend", "Python 3.12, FastAPI, Pydantic v2, Uvicorn", "Typed endpoints and automatic API documentation at /docs"],
    ["Database", "MongoDB 8.3 as a one-node replica set (rs0)", "Transactions need a replica set; the same code runs on Atlas"],
    ["Driver", "PyMongo 4", "Official driver with sessions, transactions and aggregation"],
    ["Embeddings", "TF-IDF followed by TruncatedSVD (LSA), 256 dimensions", "Works without downloading a model; can be replaced later"],
    ["Language model (optional)", "Claude through the Anthropic SDK", "Only rewrites retrieved text; without a key we build the answer from the text itself"],
    ["Login", "JWT tokens and PBKDF2-SHA256 password hashes", "Separate Citizen and Admin sign-in; the server checks the role"],
    ["Testing", "pytest, Hypothesis, FastAPI TestClient", "Logic tables, random property tests, sample profiles and API tests"],
], widths=[3.0, 6.4, 7.0], size=9)

# ---------------------------------------------------------------- 5. Data model
doc.add_heading("5. Database design and data modelling", 1)
para("We followed one rule of thumb while designing the collections: **keep together what is read together, keep "
     "separate what changes on its own, and copy anything that must never change.** In total the database has fifteen "
     "collections.")
figure(IMG / "data_model.png", "Figure 2 — Collections and how they relate")
table(["Collection", "What it stores", "Pattern", "Main indexes"], [
    ["schemes", "One scheme with its text sections, tags and current rules", "Embedded rule tree", "slug (unique); {status, level, state}; tags; text index"],
    ["scheme_versions", "A frozen copy of each published rule set", "Insert-only, versioned", "{scheme_slug, version} unique"],
    ["evaluations", "One record per scheme per check, with copies of the inputs and a checksum", "Embedded copies, insert-only", "{user_id, evaluated_at}; run_id; {scheme_slug, scheme_version}"],
    ["document_chunks", "Pieces of scheme text with their 256-number embeddings", "Linked by slug; filter fields copied in", "{scheme_slug, chunk_index} unique; {level, state}; text index"],
    ["users", "Email, password hash and role", "Plain", "email unique"],
    ["user_profiles", "A user's saved answers and consent", "Flexible sub-document", "user_id unique"],
    ["attribute_registry", "The list of questions a rule is allowed to use", "Reference data", "key unique"],
    ["ingestion_jobs", "Where each scheme is in the pipeline, and its proposed rules", "Status field", "{scheme_slug, source_hash} unique; {stage, updated_at}"],
    ["chat_logs", "Questions asked and the sources used", "Expires", "TTL on created_at (30 days)"],
    ["rate_limits", "Request counts per IP address", "Expires", "TTL on expires_at; {key, window} unique"],
    ["grievances", "One complaint with its ticket, status, category and a timeline of every update", "Embedded timeline, grows with $push", "ticket unique; {user_id, created_at}; {status, updated_at}; scheme_slug"],
    ["counters", "The last ticket number used each year", "Single counter document", "_id per year"],
    ["applications", "A citizen's saved scheme: own status, deadline, document checklist, status history", "Embedded arrays", "{user_id, scheme_slug} unique; {user_id, deadline}"],
    ["search_logs", "Each search: topic words, groups, state, number of matching schemes", "Expires after 180 days", "{state, created_at}; q_norm; TTL on expires_at"],
    ["regions", "Each state / UT with a GeoJSON point", "Reference data", "2dsphere on location; name unique"],
], widths=[2.9, 5.2, 3.3, 5.0], size=8.5)
doc.add_heading("5.1 A scheme document", 2)
para("Below is a shortened version of the stored document for Atal Pension Yojana. The rules sit inside the "
     "document, and each condition remembers the sentence it was taken from.")
code('''{ "slug": "apy", "name": "Atal Pension Yojana", "level": "central", "state": null,
  "status": "published", "current_version": 1,
  "source_url": "https://www.myscheme.gov.in/schemes/apy",
  "eligibility": { "op": "AND", "children": [
      { "id": "r1", "attribute": "age", "operator": "between", "value": [18, 40],
        "source_quote": "The minimum age of joining APY is 18 years and maximum is 40 years." },
      { "op": "NOT", "children": [
          { "id": "r2", "attribute": "is_income_tax_payer", "operator": "==", "value": true } ] } ] },
  "unmapped_conditions": ["Contribution by auto-debit from a savings bank account"] }''')
para("A condition can compare values (==, !=, <, <=, >, >=), check membership (in, not_in), check a range "
     "(between) or check that a value is present (exists). Conditions are combined with AND, OR and NOT, and these "
     "groups can be nested. Anything we could not turn into a condition is listed under unmapped_conditions and "
     "shown to the user as a reminder to check it themselves.")
doc.add_heading("5.2 An evaluation (audit) record", 2)
code('''{ "run_id": "run_7f3c...", "scheme_slug": "apy", "scheme_version": 1,
  "evaluated_at": ISODate("2026-10-21T08:12:44.120Z"),
  "profile_snapshot": { "age": 29, "is_income_tax_payer": false },
  "rules_snapshot": { ...copy of version 1 of the rules... },
  "verdict": "ELIGIBLE", "reasons": [ ... ], "missing_attributes": [],
  "checksum": "sha256:9a1e..." }''')
para("Since each record carries its own copy of the rules and answers it used, we can re-run it later without "
     "looking anything up, and it stays correct even after the scheme or the user's profile changes. We only copy "
     "the answers that the scheme's rules actually need, so we do not store more personal data than necessary.")
doc.add_heading("5.3 How the design serves each query", 2)
table(["What the app needs to do", "How the data model handles it"], [
    ["Check a profile against every scheme", "The {status, level, state} index picks published central schemes plus those of the user's state; each scheme is one read"],
    ["Show a scheme's page", "One document fetched by its unique slug"],
    ["Show a user's past checks, newest first", "The {user_id, evaluated_at} index returns them already sorted"],
    ["Show every result of one check", "Index on run_id"],
    ["Find checks that used a given rule version", "Index on {scheme_slug, scheme_version}"],
    ["Search by scheme name", "Text index on chunks, with the scheme name weighted five times higher"],
    ["Search by meaning", "A 256-number vector on each chunk; $vectorSearch on Atlas, cosine similarity locally"],
], widths=[6, 10.4], size=9)
doc.add_heading("5.4 A grievance document", 2)
para("A complaint and its whole history are read together, so the history is stored inside the complaint as an "
     "array. Each reply from the user or the officer is one more entry in that array.")
code('''{ "ticket": "GRV-2026-000001", "user_id": ObjectId("..."), "scheme_slug": "apy",
  "category": "application_status", "subject": "Application not processed",
  "status": "resolved", "priority": "normal", "run_id": "run_ab07...",
  "timeline": [
    { "at": ISODate("..."), "by": "user",  "action": "filed",         "status": "open" },
    { "at": ISODate("..."), "by": "admin", "action": "status_change", "status": "resolved",
      "note": "Bank confirmed your enrolment; your PRAN has been issued." } ] }''')
para("The optional run_id links the complaint to one of the user's own eligibility checks, so the officer can open "
     "the exact audit record the user is talking about.")
doc.add_heading("5.5 Planning for growth", 2)
bullets(["The evaluations collection grows fastest, since every check adds one record per scheme. Because records "
         "are only inserted, it can be sharded on {user_id, evaluated_at} or on a hashed run_id, and old records "
         "can be moved to cheaper storage.",
         "Each text chunk carries copies of the scheme's level, state and tags so searches can filter on them "
         "directly.",
         "A scheme document is only tens of kilobytes, nowhere near MongoDB's 16 MB limit.",
         "Most traffic is reading schemes, which secondary members of the replica set can serve."])

# ---------------------------------------------------------------- 6. CRUD
doc.add_heading("6. CRUD implementation", 1)
para("All database work is done with PyMongo inside the FastAPI route handlers. Data is checked twice before it is "
     "stored: once by Pydantic when the request arrives, and again by the collection's validator inside MongoDB.")
table(["Operation", "Collection", "Endpoint", "MongoDB call"], [
    ["Create", "users", "POST /api/auth/register", "insert_one; the unique email index blocks duplicates"],
    ["Create", "evaluations", "POST /api/eligibility/check", "insert_many, one record per scheme with a shared run_id"],
    ["Create", "scheme_versions", "scripts/seed.py (hand-checked rules)", "insert_one inside a transaction"],
    ["Create", "ingestion_jobs", "scripts/seed.py", "update_one with upsert, keyed on slug and source hash"],
    ["Read", "schemes", "GET /api/schemes", "find with filters, text search and paging"],
    ["Read", "schemes", "GET /api/schemes/{slug}", "find_one by slug"],
    ["Read", "evaluations", "GET /api/eligibility/runs/{run_id}", "find by run_id"],
    ["Read", "evaluations", "GET /api/admin/stats", "aggregation pipelines"],
    ["Update", "user_profiles", "PUT /api/profile", "update_one with $set and upsert"],
    ["Update", "schemes", "publish / retire", "update_one inside the publish transaction"],
    ["Delete", "user_profiles, evaluations", "DELETE /api/profile", "delete_one and delete_many (user asks to erase their data)"],
    ["Delete (soft)", "schemes", "POST /api/admin/schemes/{slug}/retire", "status set to retired so history is kept"],
    ["Create", "grievances", "POST /api/grievances", "insert_one, with a ticket number from an atomic counter"],
    ["Read", "grievances", "GET /api/grievances · /api/admin/grievances", "find by user_id; aggregation with $lookup for officers"],
    ["Update", "grievances", "POST .../comments · .../status", "update_one with $push to the timeline and $set on status"],
    ["Create", "applications", "POST /api/applications", "insert_one; the unique {user_id, scheme_slug} index blocks a second save"],
    ["Update", "applications", "PATCH /api/applications/{slug}[/documents/{id}]", "$set with arrayFilters on one checklist item; $push to history"],
    ["Delete", "applications", "DELETE .../documents/{id} · DELETE /api/applications/{slug}", "$pull one checklist item; delete_one to unsave"],
], widths=[2.0, 3.2, 5.4, 5.8], size=8.5)
doc.add_heading("6.1 Creating a user", 2)
code('''try:
    uid = db.users.insert_one({"email": body.email.lower(), "name": body.name,
                               "password_hash": hash_password(body.password),
                               "role": "user", "created_at": now()}).inserted_id
except DuplicateKeyError:
    raise HTTPException(409, "Email already registered")''')
doc.add_heading("6.2 Saving a profile (only with consent)", 2)
code('''db.user_profiles.update_one({"user_id": user["_id"]},
    {"$set": {"attributes": clean, "consent": {"store_profile": True, "given_at": now()}}},
    upsert=True)''')
doc.add_heading("6.3 Deleting a user's data", 2)
code('''db.user_profiles.delete_one({"user_id": user["_id"]})
db.evaluations.delete_many({"user_id": user["_id"]})''')
doc.add_heading("6.4 Adding to a grievance's timeline", 2)
code('''db.grievances.update_one({"ticket": ticket}, {
    "$set":  {"status": "resolved", "updated_at": now(), "resolved_at": now()},
    "$push": {"timeline": {"at": now(), "by": "admin", "action": "status_change",
                           "status": "resolved", "note": note}}})''')
para("One call changes the status and adds the history entry together, so the timeline always agrees with the "
     "current status.")
para("Our API tests confirm that the data is really saved: they create, read, change and delete records through the "
     "HTTP endpoints and then look in MongoDB directly to make sure the change happened.")

# ---------------------------------------------------------------- 7. Advanced
doc.add_heading("7. Advanced NoSQL features and query optimisation", 1)
doc.add_heading("7.1 Validation inside the database", 2)
para("Each important collection has a $jsonSchema validator set to strict mode. For example, the evaluations "
     "collection refuses any record whose verdict is not one of our three values or whose checksum does not start "
     "with \"sha256:\". The schemes collection limits level, status and rule-group values and checks the slug "
     "format; user_profiles rejects ages outside 0–120 and negative income or land; users checks the email format "
     "and role. During testing we tried inserting a record with the verdict \"MAYBE\", and MongoDB rejected it.")
doc.add_heading("7.2 Publishing rules inside a transaction", 2)
code('''with client.start_session() as s:
    s.with_transaction(lambda sess: (
        db.scheme_versions.insert_one(new_version, session=sess),
        db.schemes.update_one({"slug": slug},
            {"$set": {"eligibility": tree, "current_version": n + 1}}, session=sess)))''')
para("Both writes succeed together or not at all, so a scheme can never point to a version that was not saved. Each "
     "new version also records who checked it and a short note on what changed.")
doc.add_heading("7.3 Aggregation pipelines", 2)
para("The admin dashboard is built from aggregation pipelines. Two of them are shown below. The first counts "
     "verdicts for each scheme. The second finds which missing answers held back the most results, which tells us "
     "which question to ask users first.")
code('''db.evaluations.aggregate([
  {"$group": {"_id": {"scheme": "$scheme_slug", "verdict": "$verdict"}, "n": {"$sum": 1}}},
  {"$group": {"_id": "$_id.scheme", "total": {"$sum": "$n"},
              "verdicts": {"$push": {"k": "$_id.verdict", "v": "$n"}}}},
  {"$sort": {"total": -1}}, {"$limit": 10}])

db.evaluations.aggregate([
  {"$match": {"verdict": "NEEDS_MORE_INFO"}}, {"$unwind": "$missing_attributes"},
  {"$group": {"_id": "$missing_attributes", "n": {"$sum": 1}}}, {"$sort": {"n": -1}}])''')
para("A third pipeline groups schemes by state and uses $cond to count how many in each state already have "
     "published rules.")
doc.add_heading("7.4 Indexes and query performance", 2)
para("We checked our main queries with explain(\"executionStats\"). To see what each index was worth, we ran the "
     "same query a second time with hint({$natural: 1}), which makes MongoDB scan the whole collection. The numbers "
     "below come from our own dataset after a load test.")
bench = r("bench", [])
if bench:
    table(["Query", "Full scan", "Using the index", "Plan"],
          [[b["query"], f"{b['no_index']['docs_examined']} docs · {b['no_index']['ms']} ms",
            f"{b['with_index']['docs_examined']} docs · {b['with_index']['ms']} ms", b["with_index"]["stage"]] for b in bench],
          widths=[5.6, 3.4, 3.4, 4.0], size=8.5)
bullets(["Compound indexes put equality fields first and the sort field after them, so a user's history comes back "
         "in order without a separate sort step.",
         "Large fields such as embeddings are left out of query results with projections.",
         "All the records from one check are written in a single insert_many call.",
         "Giving the scheme name a higher weight in the text index makes questions that name a scheme find it first."])
doc.add_heading("7.5 TTL indexes", 2)
code('''db.chat_logs.create_index("created_at", expireAfterSeconds=30 * 24 * 3600)
db.rate_limits.create_index("expires_at", expireAfterSeconds=0)''')
para("Chat questions are removed automatically after 30 days, and each rate-limit counter removes itself once its "
     "time window is over. Nothing in our own code has to clean these up.")
doc.add_heading("7.6 Features used by the grievance portal", 2)
bullets(["**Atomic ticket numbers.** Each new complaint calls find_one_and_update on the counters collection with "
         "$inc and upsert. MongoDB does the increment in one step, so two complaints filed at the same moment can "
         "never get the same number. Our tests file two complaints in a row and check that the numbers follow on.",
         "**Embedded, growing timeline.** Replies are added with $push, so reading a complaint returns its full history "
         "in one fetch.",
         "**$lookup for the officer's view.** The admin queue uses an aggregation that joins each complaint with the "
         "filer's email from the users collection. The user's details stay in one place and are not copied into every "
         "complaint.",
         "**$facet for the dashboard.** One aggregation returns three results at once: complaints by status, complaints "
         "by category, and the average time to resolve (resolved_at minus created_at, in hours).",
         "**A validator on grievances.** The database accepts only our seven categories and four statuses, a ticket "
         "number in the GRV-YYYY-NNNNNN format, and a timeline with at least one entry."])
code('''db.counters.find_one_and_update({"_id": "grievance-2026"}, {"$inc": {"seq": 1}},
                                upsert=True, return_document=True)      # -> GRV-2026-000042

db.grievances.aggregate([{"$facet": {
    "by_status":   [{"$group": {"_id": "$status", "n": {"$sum": 1}}}],
    "by_category": [{"$group": {"_id": "$category", "n": {"$sum": 1}}}],
    "resolution":  [{"$match": {"resolved_at": {"$ne": None}}},
                    {"$group": {"_id": None, "avg_hours": {"$avg": {"$divide":
                        [{"$subtract": ["$resolved_at", "$created_at"]}, 3600000]}}}}]}}])''')
doc.add_heading("7.7 Combining text and vector search", 2)
para("For keyword matching we use MongoDB's $text operator and sort by its relevance score. For matching by meaning, "
     "every chunk stores a 256-number embedding. On Atlas the code sends a $vectorSearch query with level and state "
     "filters; on our local server it loads the embeddings from MongoDB once and computes cosine similarity with "
     "NumPy. We then merge the two ranked lists with Reciprocal Rank Fusion, where a chunk's score is the sum of "
     "1 / (60 + its rank) over both lists.")
doc.add_heading("7.8 Application tracker: updating inside embedded arrays", 2)
para("Each saved scheme is one applications document with the checklist and status history inside it. Ticking one "
     "document uses arrayFilters, so MongoDB changes exactly the matching item; adding uses $push and removing uses "
     "$pull. Reminders come from one aggregation that works out the days left with $dateDiff and splits the results "
     "with $facet. The status here is labelled as the citizen's own record, not the official government status.")
code('''db.applications.update_one({"user_id": uid, "scheme_slug": "aaby"},
    {"$set": {"documents.$[d].done": True}}, array_filters=[{"d.id": doc_id}])
db.applications.aggregate([{"$match": {"user_id": uid, "deadline": {"$ne": None}}},
    {"$addFields": {"days_left": {"$dateDiff": {"startDate": "$$NOW", "endDate": "$deadline", "unit": "day"}}}},
    {"$facet": {"due_soon": [{"$match": {"$expr": {"$lte": ["$days_left", "$reminder_days"]}}}],
                "overdue":  [{"$match": {"days_left": {"$lt": 0}}}]}}])''')
doc.add_heading("7.9 Recommendation ranking with an aggregation pipeline", 2)
para("The rules engine first decides which schemes are possible; ineligible ones are never recommended. Schemes whose "
     "rules have not been hand-checked are pre-checked against their machine-read rules and shown with an "
     "\"unverified\" label. A pipeline then scores and sorts the candidates: 45% for the conditions met (worth less when "
     "a scheme's only rules are age, citizenship or state), 40% for how well the groups the scheme is meant for match "
     "the person's ($setIntersection), and 15% for how often users with a similar profile saved it ($lookup on "
     "applications). A scheme aimed only at someone else, such as a student scholarship for a 67-year-old, is skipped. "
     "Every result lists the parts of its score in plain words.")
doc.add_heading("7.10 Coverage gap analytics: $facet, $bucket and geospatial queries", 2)
para("Supply is the schemes tagged with each beneficiary group; demand is the searches people make, logged with "
     "their state and how many schemes matched. One $facet stage computes supply by state and group, another computes "
     "demand. $bucket sorts states into bands by number of schemes, and searches into bands by number of matches. "
     "Each state is stored as a GeoJSON point with a 2dsphere index, so the map can ask for nearby states with "
     "$geoNear. The dataset has no district-level schemes, so regions are states and union territories, and for the "
     "demo 900 searches were generated by a script and clearly flagged as simulated.")
code('''db.regions.aggregate([{"$geoNear": {"near": bihar["location"], "distanceField": "dist_m",
                                     "maxDistance": 600_000, "spherical": True}}])
db.schemes.aggregate([{"$group": {"_id": "$state", "n": {"$sum": 1}}},
    {"$bucket": {"groupBy": "$n", "boundaries": [1, 5, 10, 20, 40], "default": "40+"}}])''')

# ---------------------------------------------------------------- 8. Rules engine
doc.add_heading("8. Rules engine design", 1)
para("The rules engine (api/engine/evaluate.py) is a small recursive function that walks through a scheme's rule "
     "tree. It does not touch the database, which makes it easy to test and fast enough to check every scheme in a "
     "few milliseconds.")
figure(IMG / "rule_tree.png", "Figure 3 — Checking a rule tree with three possible results", width=14)
doc.add_heading("8.1 Three possible results", 2)
para("If a user leaves a question blank, the condition that needs it becomes UNKNOWN rather than FAIL. So instead of "
     "wrongly saying \"you are not eligible\", the app says \"tell us how much land you own and we can check\".")
table(["AND", "PASS", "FAIL", "UNKNOWN"], [["PASS", "PASS", "FAIL", "UNKNOWN"], ["FAIL", "FAIL", "FAIL", "FAIL"],
                                          ["UNKNOWN", "UNKNOWN", "FAIL", "UNKNOWN"]], widths=[3, 3, 3, 3])
para("OR works the other way round: one PASS is enough to pass, all FAILs give FAIL, and anything else is UNKNOWN. "
     "NOT swaps PASS and FAIL and leaves UNKNOWN as it is. At the end, PASS is shown as Eligible, FAIL as Not "
     "eligible, and UNKNOWN as Need more information, together with the list of missing answers. When one failed "
     "condition already settles a group, the engine drops the missing answers inside that group, so users are not "
     "asked questions that would not change anything.")
doc.add_heading("8.2 Suggesting the next question", 2)
para("After a check, the engine looks at all the Need more information results and counts how often each missing "
     "answer comes up. It then suggests the most common one, because answering it will settle the largest number "
     "of schemes.")
doc.add_heading("8.3 Turning scheme text into rules", 2)
para(f"Our extractor (api/ingest/extract.py) reads the Eligibility and Exclusions sections of each scheme and "
     f"suggests a rule tree. Every suggested condition keeps the sentence it came from. If a sentence cannot be "
     f"expressed using our list of allowed questions, it goes into unmapped_conditions so that the user can see "
     f"it. Over the {len(schemes)} schemes the extractor suggested {leaf_total} conditions, and "
     f"{valid} of the suggested rule trees passed our automatic checks. The validator (api/engine/validate.py) "
     f"makes sure that only allowed questions and operators are used, that values match the question type, that "
     f"ranges make sense, that each quoted sentence really exists in the scheme text, that the numbers in a rule "
     f"also appear in its sentence, and that no two conditions contradict each other. Only rules that a person has "
     f"checked are published: we checked {len(curated)} well-known schemes by hand. The other proposals are used only "
     f"as \"unverified\" hints in the recommendations.")
para("The validator also caught a mistake of our own. In our hand-written rules for the OBC post-matric scholarship "
     "we had typed an income limit of ₹2.5 lakh. The validator refused to publish it because the quoted sentence "
     "did not match the source, and when we checked, the actual limit was ₹1.5 lakh.")
table(["Question used in the rule", "Conditions suggested"], [[a, n] for a, n in attr_counts.most_common(10)], widths=[6, 4])

# ---------------------------------------------------------------- 9. RAG
doc.add_heading("9. Grounded retrieval-augmented Q&A", 1)
bullets(["**Splitting the text.** Each scheme is cut along its own sections (Details, Benefits, Eligibility, "
         "Exclusions, Application, Documents, FAQ) and then into pieces of at most 350 words that overlap by one "
         "sentence. Before a piece is embedded we add the scheme name and section in front of it, so that it still "
         "makes sense on its own.",
         "**Finding the right pieces.** Keyword and similarity results are merged with Reciprocal Rank Fusion and "
         "the best six pieces are kept. Results can be limited to central schemes plus those of the user's state.",
         "**Sending each question to the right place.** \"Am I eligible for APY?\" goes to the rules engine. "
         "\"Schemes for women entrepreneurs\" first finds likely schemes and then checks each one with the engine. "
         "\"How do I apply?\" is answered from the text. Questions that have nothing to do with schemes get a polite "
         "refusal.",
         "**Staying close to the sources.** Each sentence of an answer must point to one of the numbered pieces it "
         "came from. The code checks every citation, and if an answer has no valid citation it is replaced with "
         "\"I don't have reliable information on that. Please check myscheme.gov.in.\"",
         "**The model never decides eligibility.** Verdicts only ever come from the rules engine. If no API key is "
         "set, the answer is built directly from the best-matching sentences, so the feature works offline too."])
doc.add_heading("9.1 Fixing wrong answers found in testing", 2)
para("While testing the app ourselves we found that some short questions got answers about the wrong scheme. For "
     "example \"car loan\" returned text about education loans, and \"atal yojna\" and \"aadhar card\" were refused as "
     "off-topic. We traced this to three causes and fixed each one.")
bullets(["**The relevance check was too loose.** A question word counted as covered if it appeared anywhere in the "
         "results put together, so any passage containing \"loan\" passed for \"car loan\". Now each passage must "
         "contain the words that say what the question is about. For questions of one or two such words, all of them "
         "must appear; for longer questions, most of them.",
         "**Misspellings were treated as off-topic.** The topic filter only knew correctly spelled words. We added a "
         "small spelling map (yojna to yojana, aadhar to aadhaar, scholership to scholarship, and others) and simple "
         "plural matching (fishermen and fisherman count as the same word). A question is now also accepted if it "
         "matches the name of a scheme in our database.",
         "**Amounts were split in half.** The sentence splitter treated the full stop in \"Rs 10.00 lakh\" as the end of a "
         "sentence. It now ignores a full stop that is directly followed by a digit."])
para("We also changed what the user sees when a question is short or has no match. A short topic such as \"atal "
     "yojna\" now returns a list of the matching schemes with a one-line description and a source link for each. When "
     "nothing in our data matches, the app says so plainly and lists schemes whose names share a word with the "
     "question, instead of answering about something else. A question that starts with \"How do I apply\" is now "
     "treated as a request for the application steps rather than an eligibility check.")
table(["Question", "Before the fix", "After the fix"], [
    ["car loan", "Sentences about education loans", "\"I couldn't find a scheme about car loan\", plus related loan schemes"],
    ["atal yojna", "Refused as off-topic", "List of the Atal schemes with descriptions and sources"],
    ["aadhar card", "Refused as off-topic", "Document lists that mention the Aadhaar card, with sources"],
    ["How do I apply for Aam Aadmi Bima Yojana?", "An eligibility verdict", "The application steps from the scheme page"],
], widths=[4.4, 5.0, 7.0], size=9)

# ---------------------------------------------------------------- 10. Integration
doc.add_heading("10. Application integration and functionality", 1)
table(["Method", "Path", "What it does", "Who can call it"], [
    ["GET", "/api/health", "Checks the server, the database and the replica set", "Anyone"],
    ["GET", "/api/attributes", "Returns the list of profile questions for the form", "Anyone"],
    ["GET", "/api/schemes", "Lists, filters and searches schemes page by page", "Anyone"],
    ["GET", "/api/schemes/{slug}", "One scheme with its rules, versions and how recently it was verified", "Anyone"],
    ["POST", "/api/eligibility/check", "Checks a profile against all schemes and saves the audit records", "Anyone"],
    ["GET", "/api/eligibility/runs/{run_id}", "Results of an earlier check", "Owner or admin"],
    ["GET", "/api/eligibility/history", "The logged-in user's past checks", "User"],
    ["POST", "/api/auth/register · /api/auth/login", "Create a citizen account, or log in on the Citizen or Admin tab", "Anyone"],
    ["GET · PUT · DELETE", "/api/profile", "Read, save (with consent) or delete the profile and history", "User"],
    ["POST", "/api/chat", "Answers a question and streams the reply with sources", "Anyone, rate-limited"],
    ["POST", "/api/recommendations", "Ranked schemes for a profile, with the reasons for each rank", "Anyone"],
    ["GET · POST · PATCH · DELETE", "/api/applications[/…]", "Save schemes, set status and deadline, tick documents; reminders", "User"],
    ["POST", "/api/admin/schemes/{slug}/retire", "Retire a scheme without deleting it", "Admin"],
    ["GET", "/api/admin/stats · /api/admin/coverage[/nearby]", "Dashboard numbers; coverage gaps; nearby states", "Admin"],
    ["POST", "/api/grievances", "File a complaint about a scheme, optionally linked to one of your checks", "User"],
    ["GET", "/api/grievances · /api/grievances/{ticket}", "Your complaints, or one complaint with its timeline", "Owner or admin"],
    ["POST", "/api/grievances/{ticket}/comments", "Add a comment, or reopen a resolved complaint", "Owner or admin"],
    ["GET", "/api/admin/grievances", "All complaints, filtered by status, with the filer's email", "Admin"],
    ["POST", "/api/admin/grievances/{ticket}/status", "Change status and priority, with a note to the citizen", "Admin"],
], widths=[2.4, 5.6, 6.2, 2.2], size=8.5)
para("The **public site** is one long page that scrolls smoothly between sections. **Check** shows the questions "
     "in four groups across the page, then the ranked recommendations and the results with a reason for each "
     "condition. **Ask** streams answers with source links. **Schemes** opens a page with the rule tree and its "
     "versions. **Grievances** lets a citizen file and follow a complaint. **My applications** holds saved schemes, "
     "their checklist and reminders, and **My account** shows past checks and a button to delete one's data. Login "
     "is a **separate page** with Citizen and Admin tabs. The **officers' dashboard** has overview numbers, the "
     "coverage-gap map and tables, and the grievance queue.")
shots = sorted((IMG).glob("screen_*.png"))
caps = {p: f"Figure {3 + i} — {p.stem.split('_', 2)[2].replace('_', ' ').capitalize()}" for i, p in enumerate(shots, 1)}
wide = [p for p in shots if "recommendations" in p.stem or "coverage" in p.stem]  # the two most detailed screens
small = [p for p in shots if p not in wide]
figure_grid(small, [caps[p] for p in small])
for p in wide:
    figure(p, caps[p], width=15.5)

# ---------------------------------------------------------------- 11. Testing & security
doc.add_heading("11. Testing, security and data validation", 1)
doc.add_heading("11.1 Checking input at three levels", 2)
bullets(["**At the API**, Pydantic models enforce ranges (age 0–120, income zero or more, marks 0–100), allowed "
         "values taken from the question list, and length limits. A profile field we do not recognise is rejected.",
         "**Before publishing**, the rule validator from section 8.3 runs.",
         "**In the database**, strict JSON Schema validators and unique indexes act as the final safety net."])
doc.add_heading("11.2 Login, roles and privacy", 2)
bullets(["Passwords are stored as PBKDF2-SHA256 hashes (240,000 rounds, a random salt per user) and compared in "
         "constant time.",
         "After login the user receives a JWT that expires after 12 hours; the signing secret is read from an "
         "environment variable.",
         "Anyone can check eligibility without an account, and their answers then stay in their own browser. Logged-in "
         "users can save a profile and see their history. Citizens and officers sign in on separate tabs; the server "
         "rejects a citizen account on the Admin tab, and only admins can open /api/admin/* (HTTP 403 otherwise).",
         "Audit records are only ever inserted by the API, and each one has a SHA-256 checksum. In a test we changed "
         "one stored value by hand and re-checking the record reported the checksum as invalid.",
         "We never ask for Aadhaar or any other ID number. A profile is saved only after the user agrees, the user "
         "can delete it together with their history at any time, and chat logs expire on their own.",
         "The chat endpoint is limited per IP address, using counters kept in a collection with a TTL index.",
         "A grievance can only be read or commented on by the user who filed it or by an admin; anyone else gets "
         "HTTP 403. Filing one requires login, a linked check must be from the user's own history, and the form asks "
         "users not to type Aadhaar or bank numbers.",
         "Queries are built as Python dictionaries rather than pasted strings, and user input is never used as a "
         "query operator, which protects against injection."])
doc.add_heading("11.3 Automated tests", 2)
table(["Test file", "What it checks"], [
    ["tests/test_engine.py", "AND, OR and NOT for every mix of PASS, FAIL and UNKNOWN; dropping questions that no longer matter; the suggested question"],
    ["tests/test_properties.py", "Hypothesis generates random rule trees and checks that NOT(NOT(x)) equals x and that answering a question never undoes a settled result"],
    ["tests/test_fixtures.py", "Sample profiles for the hand-checked schemes give the expected verdict"],
    ["tests/test_validate.py", "Unknown questions or operators, bad values, made-up quotes, numbers missing from quotes and contradictions are all rejected"],
    ["tests/test_api.py", "CRUD through HTTP on a test database; citizen vs admin login; checksum tamper check; validators; grievance flow; tracker (arrayFilters, $push / $pull, reminders); ranking (no ineligible schemes, audience check); coverage analytics and $geoNear"],
], widths=[5.2, 11.2], size=9)
tests = r("tests")
if tests != "—":
    para(f"On the last run, **{tests['passed']} tests passed and {tests['failed']} failed**, taking {tests['seconds']} "
         f"seconds in total.")

# ---------------------------------------------------------------- 12. Results
doc.add_heading("12. Results", 1)
table(["Measure", "Value"], [
    ["Scheme PDFs downloaded (duplicates removed)", n_pdfs],
    ["Schemes parsed that have an Eligibility section", len(schemes)],
    ["Central / state schemes", f"{levels['central']} / {levels['state']}"],
    ["States and union territories covered", n_states],
    ["Conditions suggested by the extractor", leaf_total],
    ["Suggested rule trees that passed automatic checks", f"{valid} of {len(schemes)}"],
    ["Schemes with hand-checked, published rules", len(curated)],
    ["Schemes published in total", r("published")],
    ["Text pieces indexed for search", r("chunks")],
    ["Audit records in the database during benchmarking", r("evaluations")],
    ["Automated tests", r("tests.summary")],
    ["Sample profiles that got the expected verdict", r("fixture_accuracy")],
    ["Median time to check all published schemes and save the records", r("check_ms")],
], widths=[11, 5.4])
ev = r("rag_eval")
if ev != "—":
    para(f"To test the search we used {ev['n']} questions where we already knew which scheme held the answer, and "
         f"measured how often that scheme appeared in the top five results (Recall@5) and how high it was ranked "
         f"(MRR).")
    table(["Setup", "Recall@5", "MRR"], [[k, f"{v['recall@5']:.2f}", f"{v['mrr']:.2f}"] for k, v in ev["configs"].items()],
          widths=[6, 4, 4])
    c = ev["configs"]
    para(f"\"Named\" questions mention the scheme by name, while \"descriptive\" ones describe what the scheme does "
         f"without naming it. Keyword search was almost perfect when the name was there ({c['keyword · named']['recall@5']:.2f}) "
         f"but poor otherwise ({c['keyword · descriptive']['recall@5']:.2f}); similarity search behaved the other way "
         f"round. The combined search kept the strong result on named questions ({c['hybrid · named']['recall@5']:.2f}), "
         f"but on descriptive questions it scored {c['hybrid · descriptive']['recall@5']:.2f}, which is lower than "
         f"similarity search alone ({c['vector · descriptive']['recall@5']:.2f}). Weak keyword matches pull the merged "
         f"ranking down. We are reporting this as we found it; the fix we plan to try is to give the two lists "
         f"different weights, or to use keyword search only when a scheme name is detected. Out of 10 questions that "
         f"had nothing to do with schemes or had no answer in our data, {ev['refusal'].split('/')[0]} were correctly "
         f"refused.")

doc.add_heading("12.1 What the new features showed", 2)
from api.analytics import coverage_report  # noqa: E402
from api.db import get_db  # noqa: E402
from api.recommend import recommend  # noqa: E402
_db = get_db()
_people = [
    ("Woman farmer, 45, Punjab, BPL", {"age": 45, "gender": "female", "state": "Punjab", "category": "General", "annual_income": 50000,
                                        "occupation": "farmer", "is_student": False, "is_bpl": True, "is_income_tax_payer": True, "is_citizen": True}),
    ("SC student, 19, Bihar", {"age": 19, "gender": "male", "state": "Bihar", "category": "SC", "annual_income": 120000, "occupation": "student",
                               "is_student": True, "education_level": "class_12", "marks_percent": 72, "is_citizen": True, "is_income_tax_payer": False}),
    ("Person with disability, 24, Tamil Nadu", {"age": 24, "gender": "male", "state": "Tamil Nadu", "has_disability": True, "disability_percent": 60,
                                                "occupation": "unemployed", "is_citizen": True, "is_income_tax_payer": False}),
    ("Woman weaver, 29, Assam, BPL", {"age": 29, "state": "Assam", "annual_income": 90000, "is_bpl": True, "occupation": "weaver",
                                      "gender": "female", "is_citizen": True, "is_income_tax_payer": False}),
]
para("We ran the recommendation ranking for four different people. Each gets a different list, led by schemes meant "
     "for their own groups; \"(unverified)\" marks schemes whose rules have not been hand-checked.")
_rows = []
for who, prof in _people:
    items = recommend(_db, prof, limit=3)["items"]
    _rows.append([who, "; ".join(i["name"] + ("" if i["verified"] else " (unverified)") for i in items)])
table(["Person", "Top three recommendations"], _rows, widths=[5.0, 11.4], size=8.5)
_cov = coverage_report(_db)
para(f"The coverage dashboard analysed {_cov['searches']} searches ({_cov['demo_searches']} of them simulated for the demo). "
     f"The regions with the largest gap between searches and state schemes were:")
table(["State / UT", "State schemes", "Searches", "Poorly served", "Gap index"],
      [[x["name"], x["state_schemes"], x["searches"], f"{round(x['low_match_rate'] * 100)}%", f"{x['gap_index']:.2f}"]
       for x in _cov["regions"][:5]], widths=[4.6, 2.8, 2.6, 3.0, 3.4], size=8.5)

# ---------------------------------------------------------------- 13. Limitations
doc.add_heading("13. Limitations and future work", 1)
bullets(["Our data is a snapshot of pages from 2024, so some details may be out of date. The app shows when each "
         "scheme was last verified and warns when that was more than 90 days ago. A real deployment should pull data "
         "from the official API Setu myScheme service.",
         "The extractor works with text patterns. A language model with a fixed output format could replace it while "
         "keeping the same validator and human review. Some conditions, such as \"no family member in government "
         "service\", still cannot be expressed and are shown as warnings.",
         "Our LSA embeddings only handle English well. A multilingual model would let someone ask in Hindi and still "
         "find the English text.",
         "The grievance portal and the tracker do not yet send email or SMS; users see updates when they log in.",
         "Recommendations from unreviewed schemes rely on machine-read rules and can miss details, which is why they "
         "carry an \"unverified\" label. Coverage analytics are by state, and the demo demand data is simulated.",
         "Next, we would like to move to MongoDB Atlas to use its built-in vector and text search indexes, add a Hindi "
         "interface, and build a test set that measures how faithful the answers are to their sources."])

# ---------------------------------------------------------------- 14. Conclusion
doc.add_heading("14. Conclusion", 1)
para("Working on this project showed us that MongoDB's document model suits government scheme data well. The data "
     "has many conditions and varies a lot from one scheme to another. Each scheme fits in one document along with "
     "its nested rules. Rule changes become new versions saved in a transaction, and every eligibility decision is a "
     "single record, with its own checksum, that can be re-checked later without any joins. Validators, unique and "
     "compound indexes, text and vector search, aggregation pipelines and TTL indexes together covered our needs "
     "for correct data, speed and privacy. On top of the database, the three-result rules engine gives answers that "
     "users can understand, the Q&A feature answers with sources but leaves the eligibility decision to the "
     "engine. Ranked recommendations, the application tracker, the grievance portal and the coverage-gap dashboard "
     "show how far aggregation pipelines, embedded arrays and geospatial indexes can carry an application.")

doc.add_heading("15. References", 1)
bullets(["MongoDB Manual: data modelling, schema validation, transactions, TTL indexes, text indexes and aggregation. https://www.mongodb.com/docs/manual/",
         "MongoDB Atlas Vector Search and Atlas Search documentation. https://www.mongodb.com/docs/atlas/",
         "myScheme, the national platform for government schemes. https://www.myscheme.gov.in/",
         "shrijayan/gov_myscheme dataset on Hugging Face. https://huggingface.co/datasets/shrijayan/gov_myscheme",
         "FastAPI documentation, https://fastapi.tiangolo.com/, and PyMongo documentation, https://pymongo.readthedocs.io/",
         "Cormack, G. V., Clarke, C. L. A. and Büttcher, S. (2009). Reciprocal Rank Fusion outperforms Condorcet and individual rank learning methods. SIGIR.",
         "Lewis, P. et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS."], style="List Number")

doc.add_heading("Appendix A — Setup and run instructions", 1)
code('''pip install -r requirements.txt
powershell -ExecutionPolicy Bypass -File scripts\\start_mongo.ps1
$env:HF_TOKEN="<your token>"; python scripts\\download_hf.py
python scripts\\seed.py
python scripts\\simulate_searches.py
uvicorn api.main:app --port 8000
pytest -q''')
para("The commands, in order: install the Python packages; start MongoDB as a one-node replica set (needed for "
     "transactions); download the scheme PDFs; parse them, suggest and publish rules, and build the search index; "
     "add flagged demo searches for the coverage dashboard; "
     "start the app at http://127.0.0.1:8000 (API documentation at /docs); and run the tests. The seed script also "
     "creates two demo accounts, admin@example.org and user@example.org, with random passwords saved to the "
     "git-ignored file .env.demo for local use only.", size=10)

# ---------------------------------------------------------------- Appendix B: plagiarism report
plag_pages = sorted(IMG.glob("plagiarism_report_p*.png"))
if plag_pages:
    page_break()
    doc.add_heading("Appendix B — Plagiarism report", 1)
    para("The similarity check for this report is attached below, as generated by the plagiarism-checking tool.", size=10)
    for i, pg in enumerate(plag_pages):
        if i:
            page_break()
        doc.add_picture(str(pg), height=Cm(22.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print("wrote", OUT)
